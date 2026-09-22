"""
Builds excel/funnel_financial_model.xlsx from live PostgreSQL data.

Produces a workbook with real Excel formulas (not pasted static values):
  - Monthly Data (raw pulled from Postgres)
  - Funnel Model: won-rate %, activation-rate %, lead-to-activated %, a chart
  - Activation Opportunity What-If: revenue opportunity from closing the
    won-to-activated gap, with adjustable input cells
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment
from sqlalchemy import create_engine

ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "excel" / "funnel_financial_model.xlsx"

ENGINE = create_engine("postgresql+psycopg2://postgres:postgres@localhost:5432/merchant_funnel_analytics")

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
INPUT_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")


def style_header(ws, row: int, n_cols: int) -> None:
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def autosize(ws, widths: dict) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def fetch_monthly() -> pd.DataFrame:
    q = """
        SELECT DATE_TRUNC('month', mql.first_contact_date)::date AS month,
               COUNT(DISTINCT mql.mql_id) AS leads,
               COUNT(DISTINCT cd.mql_id) AS won,
               COUNT(DISTINCT sr.seller_id) AS activated,
               COALESCE(SUM(sr.total_revenue), 0) AS revenue
        FROM marketing_qualified_leads mql
        LEFT JOIN closed_deals cd ON cd.mql_id = mql.mql_id
        LEFT JOIN seller_revenue sr ON sr.seller_id = cd.seller_id
        WHERE mql.first_contact_date < '2018-06-01'
        GROUP BY 1 ORDER BY 1
    """
    return pd.read_sql(q, ENGINE)


def fetch_totals() -> dict:
    q = """
        SELECT
            (SELECT COUNT(*) FROM closed_deals) AS won_deals,
            (SELECT COUNT(DISTINCT sr.seller_id) FROM closed_deals cd
                JOIN seller_revenue sr ON sr.seller_id = cd.seller_id) AS activated_sellers,
            (SELECT ROUND(AVG(sr.total_revenue), 2) FROM closed_deals cd
                JOIN seller_revenue sr ON sr.seller_id = cd.seller_id) AS avg_revenue_per_activated_seller
    """
    return pd.read_sql(q, ENGINE).iloc[0].to_dict()


def build_monthly_sheet(wb: Workbook, monthly: pd.DataFrame) -> None:
    ws = wb.active
    ws.title = "Monthly Data"
    headers = ["Month", "Leads", "Won Deals", "Activated Sellers", "Revenue (R$)"]
    ws.append(headers)
    style_header(ws, 1, len(headers))
    for _, r in monthly.iterrows():
        ws.append([r["month"], int(r["leads"]), int(r["won"]), int(r["activated"]), round(r["revenue"], 2)])
    autosize(ws, {"A": 14, "B": 10, "C": 12, "D": 16, "E": 16})
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=1).number_format = "yyyy-mm"


def build_funnel_model_sheet(wb: Workbook, n_months: int) -> None:
    ws = wb.create_sheet("Funnel Model")
    headers = ["Month", "Leads", "Won Rate %", "Activation Rate %", "Lead-to-Activated %", "Revenue (R$)"]
    ws.append(headers)
    style_header(ws, 1, len(headers))

    src = "'Monthly Data'!"
    for i in range(n_months):
        r = i + 2
        ws.cell(row=r, column=1, value=f"={src}A{r}")
        ws.cell(row=r, column=2, value=f"={src}B{r}")
        ws.cell(row=r, column=3, value=f"=IFERROR({src}C{r}/{src}B{r},0)")
        ws.cell(row=r, column=4, value=f"=IFERROR({src}D{r}/{src}C{r},0)")
        ws.cell(row=r, column=5, value=f"=IFERROR({src}D{r}/{src}B{r},0)")
        ws.cell(row=r, column=6, value=f"={src}E{r}")
        for c in (3, 4, 5):
            ws.cell(row=r, column=c).number_format = "0.0%"
        ws.cell(row=r, column=1).number_format = "yyyy-mm"
        ws.cell(row=r, column=6).number_format = "#,##0.00"

    autosize(ws, {"A": 14, "B": 10, "C": 12, "D": 16, "E": 18, "F": 16})

    chart = LineChart()
    chart.title = "Won Rate vs Activation Rate Over Time"
    chart.y_axis.title = "Rate"
    chart.x_axis.title = "Month"
    data = Reference(ws, min_col=3, max_col=4, min_row=1, max_row=n_months + 1)
    chart.add_data(data, titles_from_data=True)
    cats = Reference(ws, min_col=1, min_row=2, max_row=n_months + 1)
    chart.set_categories(cats)
    chart.width = 22
    chart.height = 10
    ws.add_chart(chart, "H2")


def build_whatif_sheet(wb: Workbook, totals: dict) -> None:
    ws = wb.create_sheet("Activation Opportunity")
    ws["A1"] = "Won-to-Activated Revenue Opportunity Model"
    ws["A1"].font = Font(bold=True, size=13)

    rows = [
        ("Won deals (actual)", int(totals["won_deals"])),
        ("Activated sellers (actual)", int(totals["activated_sellers"])),
        ("Current activation rate", None),
        ("Avg revenue per activated seller (R$)", round(float(totals["avg_revenue_per_activated_seller"]), 2)),
        ("", None),
        ("--- Adjustable Inputs ---", None),
        ("Target activation rate", 0.60),
    ]
    r = 3
    for label, val in rows:
        ws.cell(row=r, column=1, value=label)
        if val is not None:
            cell = ws.cell(row=r, column=2, value=val)
            if "rate" in label.lower():
                cell.number_format = "0.0%"
                cell.fill = INPUT_FILL
        r += 1

    ws["B5"] = "=B4/B3"
    ws["B5"].number_format = "0.0%"

    r += 1
    ws.cell(row=r, column=1, value="--- Model Output ---").font = Font(bold=True)
    out_start = r + 1
    outputs = [
        ("Additional sellers activated to hit target", "=MAX(ROUND(B9*B3,0)-B4,0)"),
        ("Additional revenue from closing the activation gap (R$)", f"=B{out_start}*B6"),
    ]
    r = out_start
    for label, formula in outputs:
        ws.cell(row=r, column=1, value=label)
        cell = ws.cell(row=r, column=2, value=formula)
        cell.number_format = "#,##0" if "sellers" in label.lower() else "#,##0.00"
        r += 1

    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 18

    note_row = r + 2
    ws.cell(row=note_row, column=1,
            value=("Model logic: additional revenue = (sellers needed to reach target activation rate) x "
                   "(avg revenue per activated seller). Change the yellow target-rate cell to re-run the "
                   "scenario live -- e.g. set it to 100% to see the full size of the won-but-never-activated gap."))
    ws.cell(row=note_row, column=1).alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row + 2, end_column=6)


def main() -> None:
    monthly = fetch_monthly()
    totals = fetch_totals()

    wb = Workbook()
    build_monthly_sheet(wb, monthly)
    build_funnel_model_sheet(wb, len(monthly))
    build_whatif_sheet(wb, totals)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_PATH)
    print(f"Workbook written to {OUT_PATH}")


if __name__ == "__main__":
    main()
